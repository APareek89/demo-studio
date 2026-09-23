"""Bounded, model-scoped public-source discovery. No hosted crawler.

Every network hop is checked and connected to its checked IP (including redirects).
Coverage is an audit of the eligible frontier, never a promise of whole-site coverage.
"""
from __future__ import annotations

import hashlib
import ipaddress
import os
from pathlib import Path
import shutil
import re
import socket
import time
from collections import deque
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeout
from urllib.parse import urljoin, urlsplit, urlunsplit
from urllib.robotparser import RobotFileParser
from xml.etree import ElementTree

import httpx
from bs4 import BeautifulSoup

from . import store

UA = "DemoStudio/0.2 (product evidence reader)"
_RESOLVER = ThreadPoolExecutor(max_workers=4, thread_name_prefix="source-dns")
DEFAULT_BUDGET = {"html_pages": 40, "documents": 10, "document_pages": 300, "depth": 4, "seconds": 180}
DOCUMENT = re.compile(r"\.(pdf|docx?)(?:$|[?#])", re.I)
TOPICS = ("specification", "variant", "exterior", "interior", "safety", "technology", "ownership", "warranty", "faq", "price", "brochure")
EXCLUDE = re.compile(r"/(?:news|press|career|dealer|blog|investor|events?)(?:/|[-?])", re.I)


def canonical_url(url: str) -> str:
    p = urlsplit(url)
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path or "/", p.query, ""))


def _public_address(url: str, allowed_hosts=None, *, timeout: float = 5) -> tuple[str, str]:
    p = urlsplit(url)
    if p.scheme not in {"http", "https"} or not p.hostname or p.username or p.password:
        raise ValueError("Only public HTTP(S) URLs without credentials are supported")
    if p.port not in {None, 80, 443}:
        raise ValueError("Only public web ports 80 and 443 are supported")
    host = p.hostname.encode("idna").decode("ascii")
    if allowed_hosts is not None and host.lower() not in {h.lower() for h in allowed_hosts}:
        raise ValueError("URL host is outside the allowed source scope")
    lookup = _RESOLVER.submit(socket.getaddrinfo, host, p.port or (443 if p.scheme == "https" else 80), type=socket.SOCK_STREAM)
    try:
        answers = lookup.result(timeout=max(.1, timeout))
    except FutureTimeout as exc:
        lookup.cancel()
        raise TimeoutError("Source DNS lookup exceeded its time budget") from exc
    addresses = sorted({a[4][0] for a in answers})
    if not addresses or any(not ipaddress.ip_address(a).is_global or ipaddress.ip_address(a).is_multicast or ipaddress.ip_address(a).is_reserved for a in addresses):
        raise ValueError("Private, local and reserved network addresses are not permitted")
    return host, addresses[0]


def parse_html(raw: bytes | str, url: str) -> dict:
    soup = BeautifulSoup(raw, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    links = [{"url": canonical_url(urljoin(url, a.get("href", ""))), "label": a.get_text(" ", strip=True)}
             for a in soup.find_all("a", href=True) if urlsplit(urljoin(url, a["href"])).scheme in {"http", "https"}]
    images = list(dict.fromkeys(urljoin(url, i.get("src") or i.get("data-src")) for i in soup.find_all("img")
                               if (i.get("src") or i.get("data-src")) and not (i.get("src") or i.get("data-src")).startswith("data:")))[:60]
    for node in soup(["script", "style", "noscript", "svg", "iframe", "nav", "header", "footer"]):
        node.decompose()
    sections, heading, seen = [], title, set()
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table", "div", "span"]):
        if node.find_parent("table") or (node.name in {"div", "span", "li"} and node.find(["h1", "h2", "h3", "h4", "p", "li", "table", "div"])):
            continue
        if node.name.startswith("h") and node.name != "html":
            heading = node.get_text(" ", strip=True)
        if node.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in row.find_all(["th", "td"], recursive=False)] for row in node.find_all("tr")]
            text = "\n".join(" | ".join(row) for row in rows if row)
            # Keep adjacent caveats with the table; do not infer column associations.
            notes = node.find_next_sibling()
            footnote = notes.get_text(" ", strip=True) if notes and notes.name in {"p", "small"} else ""
            if footnote:
                text += "\nFootnote: " + footnote
            extra = {"rows": rows, "footnote": footnote}
        else:
            text, extra = node.get_text(" ", strip=True), {}
        if not text or text in seen:
            continue
        seen.add(text)
        sections.append({"locator": node.get("id") or f"{heading} · block {len(sections)+1}", "heading": heading,
                         "kind": "table" if node.name == "table" else "text", "text": text, **extra})
    text = "\n\n".join(s["text"] for s in sections)
    warnings = [] if len(text.strip()) >= 80 else ["Little readable HTML; authoring Read may attempt a rendered fallback. Coverage may be incomplete."]
    return {"title": title, "text": text, "sections": sections, "links": links, "images": images, "warnings": warnings}


def fetch_public(url: str, *, timeout: float = 15, max_bytes: int = 12_000_000, allowed_hosts=None) -> dict:
    """Fetch one public page/document with a wall-clock bound and exact evidence text.

    Returns bytes for immutable ingestion, plus paragraphs/tables for runtime tools.
    Callers must treat returned web content as untrusted evidence, never instructions.
    """
    start, current = time.monotonic(), canonical_url(url)
    with httpx.Client(follow_redirects=False, trust_env=False, headers={"User-Agent": UA}) as client:
        for _ in range(6):
            host, address = _public_address(current, allowed_hosts, timeout=min(5, max(.1, timeout - (time.monotonic() - start))))
            remaining = timeout - (time.monotonic() - start)
            if remaining <= 0:
                raise TimeoutError("Source fetch time budget exhausted")
            target = httpx.URL(current).copy_with(host=address)
            # Pin DNS resolution; TLS still validates the original source hostname.
            with client.stream("GET", target, headers={"Host": urlsplit(current).netloc},
                               extensions={"sni_hostname": host}, timeout=min(remaining, 15)) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    current = canonical_url(urljoin(current, response.headers.get("location", "")))
                    continue
                response.raise_for_status()
                chunks, size = [], 0
                for chunk in response.iter_bytes():
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError(f"Source exceeds the {max_bytes} byte limit")
                    if time.monotonic() - start > timeout:
                        raise TimeoutError("Source fetch time budget exhausted")
                    chunks.append(chunk)
                raw = b"".join(chunks)
                mime = response.headers.get("content-type", "").split(";")[0].lower()
                result = {"url": url, "final_url": current, "raw": raw, "mime": mime, "hash": hashlib.sha256(raw).hexdigest(),
                          "fetched_at": store.now(), "effective_date": response.headers.get("last-modified", ""), "error": None}
                if "html" in mime or (not mime and not DOCUMENT.search(current)):
                    result.update(parse_html(raw, current))
                elif mime == "application/pdf" or raw.startswith(b"%PDF"):
                    from io import BytesIO
                    from .sources import extract_pdf
                    pdf = extract_pdf(BytesIO(raw), timeout=max(.1, timeout - (time.monotonic() - start)))
                    if pdf["error"]:
                        raise ValueError(pdf["error"])
                    if any(len(p.strip()) < 30 for p in pdf["pages"]):
                        pdf["warnings"].append("Some PDF pages contain little text; local OCR is available during authoring Read only.")
                    result.update(title=current.rsplit("/", 1)[-1], text="\n\n".join(pdf["pages"]), sections=pdf["sections"], pages=pdf["pages"],
                                  page_count=pdf["page_count"], links=[], images=[], warnings=pdf["warnings"], extraction_version=2)
                elif current.lower().split("?")[0].endswith(".docx") or "wordprocessingml" in mime:
                    from io import BytesIO
                    from zipfile import ZipFile
                    from docx import Document
                    with ZipFile(BytesIO(raw)) as archive:
                        if sum(item.file_size for item in archive.infolist()) > 40_000_000:
                            raise ValueError("Expanded document exceeds the extraction byte limit")
                    document = Document(BytesIO(raw))
                    sections = [{"locator": f"paragraph {i+1}", "kind": "text", "text": p.text} for i, p in enumerate(document.paragraphs) if p.text.strip()]
                    for i, table in enumerate(document.tables):
                        rows = [[cell.text for cell in row.cells] for row in table.rows]
                        sections.append({"locator": f"table {i+1}", "kind": "table", "rows": rows, "text": "\n".join(" | ".join(row) for row in rows)})
                    result.update(title=current.rsplit("/", 1)[-1], text="\n\n".join(s["text"] for s in sections), sections=sections, links=[], images=[], warnings=[])
                elif current.lower().split("?")[0].endswith(".doc"):
                    result.update(title=current.rsplit("/", 1)[-1], text="", sections=[], links=[], images=[], warnings=["Legacy .doc extraction is unsupported; provide .docx or PDF."])
                else:
                    text = raw.decode("utf-8", errors="replace")
                    result.update(title=current.rsplit("/", 1)[-1], text=text, sections=[{"locator": "document", "kind": "text", "text": text}],
                                  links=[], images=[], warnings=[])
                return result
    raise ValueError("Too many source redirects")


def _render_executable(chromium) -> str:
    """Select an installed executable; never install browsers or weaken launch flags."""
    configured = os.environ.get("CRAWL_BROWSER_EXECUTABLE", "").strip()
    if configured:
        executable = Path(configured).expanduser()
        if not executable.is_absolute() or not executable.is_file() or not os.access(executable, os.X_OK):
            raise RuntimeError("CRAWL_BROWSER_EXECUTABLE must name an existing executable file by absolute path")
        return str(executable)
    mac_chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    if mac_chrome.is_file() and os.access(mac_chrome, os.X_OK):
        return str(mac_chrome)
    bundled = Path(chromium.executable_path)
    if bundled.is_file() and os.access(bundled, os.X_OK):
        return str(bundled)
    raise RuntimeError("Rendered extraction unavailable: no installed Chromium browser. "
                       "Configure CRAWL_BROWSER_EXECUTABLE or install Playwright Chromium and its OS dependencies for the service user.")


def render_public(url: str, *, timeout: float = 25, max_requests: int = 60, fetcher=None) -> dict:
    """Isolated local Chromium rendering, with all page HTTP fulfilled by safe fetch.

    No inherited cookies, service workers, downloads, sockets or private origins.
    This adapter is used only for authoring Read, never a customer browser session.
    """
    from playwright.sync_api import sync_playwright
    fetcher = fetcher or fetch_public
    started, requests, blocked = time.monotonic(), [], []
    allowed_hosts = {urlsplit(url).hostname}
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=_render_executable(playwright.chromium), headless=True, chromium_sandbox=True, timeout=int(timeout * 1000),
                                              args=["--disable-background-networking", "--disable-sync", "--no-first-run", "--force-webrtc-ip-handling-policy=disable_non_proxied_udp"])
        context = browser.new_context(service_workers="block", accept_downloads=False, user_agent=UA)
        context.route_web_socket("**/*", lambda route: route.close())
        context.add_init_script("window.RTCPeerConnection = undefined; window.Worker = undefined; window.SharedWorker = undefined;")
        def route_request(route):
            request = route.request
            remaining = timeout - (time.monotonic() - started)
            if (request.method != "GET" or request.resource_type in {"image", "media", "font"}
                    or remaining <= 0 or len(requests) >= max_requests or urlsplit(request.url).hostname not in allowed_hosts):
                blocked.append(request.url)
                route.abort()
                return
            requests.append(request.url)
            try:
                result = fetcher(request.url, timeout=min(10, remaining), max_bytes=8_000_000, allowed_hosts=allowed_hosts)
                route.fulfill(status=200, content_type=result.get("mime") or "text/html", body=result.get("raw", b""))
            except Exception as exc:
                blocked.append(f"{request.url}: {str(exc)[:120]}")
                route.abort()
        context.route("**/*", route_request)
        page = context.new_page()
        page.on("popup", lambda popup: popup.close())
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=max(1000, int((timeout - (time.monotonic() - started)) * 1000)))
            remaining = timeout - (time.monotonic() - started)
            if remaining > 0:
                try:
                    page.wait_for_load_state("networkidle", timeout=min(5000, int(remaining * 1000)))
                except Exception:
                    pass
            raw, final_url = page.content().encode(), page.url
            result = {"url": url, "final_url": final_url, "raw": raw, "mime": "text/html", "fetched_at": store.now(),
                      "hash": hashlib.sha256(raw).hexdigest(), "error": None, **parse_html(raw, final_url)}
            result["render"] = {"method": "local_chrome_playwright", "requests": requests, "blocked": blocked, "seconds": round(time.monotonic()-started, 3)}
            if blocked:
                result["warnings"].append(f"Renderer blocked or could not read {len(blocked)} requests; source completeness needs review.")
            return result
        finally:
            context.close()
            browser.close()


def _model_tokens(seed: dict, demo: dict) -> list[str]:
    """One canonical model slug, not arbitrary words from the demo's marketing title."""
    explicit = seed.get("scope", {}).get("model") or seed.get("crawl_model")
    brand = r"^(?:hyundai|kia|tata|maruti(?: suzuki)?|suzuki|toyota|honda|mahindra|volkswagen|skoda|renault|nissan|audi|bmw)\s+"
    if explicit:
        return [re.sub(r"[^a-z0-9]+", "-", re.sub(brand, "", explicit.casefold())).strip("-")]
    segments = [part.casefold() for part in urlsplit(seed.get("url", "")).path.split("/") if part]
    for marker in ("find-a-car", "models", "car-models", "our-vehicles", "our-models", "vehicles", "cars", "arena"):
        if marker in segments and segments.index(marker) + 1 < len(segments):
            candidate = segments[segments.index(marker) + 1].removesuffix(".html")
            if candidate not in {"all", "all-vehicles", "overview", "suv", "suvs", "index"}:
                return [candidate]
    ignore = {"www", "com", "html", "htm", "in", "en", "india", "new", "the", "all", "cars", "car", "suv", "suvs", "hyundai", "tata", "maruti", "suzuki", "demo", "studio", "product", "motor", "motors", "overview", "highlights", "specification", "specifications", "features", "price", "brochure", "pdf", "docx", "document", "content", "dam", "data", "brochures", "download", "downloads"}
    candidates = [part.removesuffix(".html") for part in segments if len(part) > 2 and part not in ignore and not DOCUMENT.search(part)]
    # A product title may contain a slogan; it can confirm a URL segment, never
    # introduce new tokens that authorize unrelated pages.
    name = (demo.get("product", {}).get("name") or demo.get("name", "")).casefold()
    if seed.get("role") != "competitor":
        matched = [part for part in candidates if re.search(r"(?<!\w)" + re.escape(part.replace("-", " ")) + r"(?!\w)", name)]
        if matched:
            return [matched[0]]
    return [candidates[0]] if len(candidates) == 1 else []


def _locale_prefix(url: str) -> tuple[str, ...]:
    segments = tuple(p.casefold() for p in urlsplit(url).path.split("/") if p)
    if len(segments) >= 2 and re.fullmatch(r"[a-z]{2}", segments[0]) and re.fullmatch(r"[a-z]{2}(?:-[a-z]{2})?", segments[1]):
        return segments[:2]
    if segments and re.fullmatch(r"[a-z]{2}(?:-[a-z]{2})?", segments[0]):
        return segments[:1]
    return ()


def _eligible(url: str, label: str, seed_url: str, tokens: list[str], *, linked: bool = True) -> bool:
    p, seed = urlsplit(url), urlsplit(seed_url)
    if EXCLUDE.search(p.path) or not tokens:
        return False
    segments = tuple(part.casefold() for part in p.path.split("/") if part)
    locale = _locale_prefix(seed_url)
    compact = lambda value: re.sub(r"[^a-z0-9]", "", value.casefold())
    models = {compact(token) for token in tokens}
    model_segment = any(compact(re.sub(r"\.(?:html?|aspx)$", "", part)) in models for part in segments)
    normalized_label = re.sub(r"[^a-z0-9]+", "-", label.casefold()).strip("-")
    named_label = any(re.search(r"(?:^|-)" + re.escape(token) + r"(?:-|$)", normalized_label) for token in tokens)
    if DOCUMENT.search(url):
        path_text = re.sub(r"[^a-z0-9]+", "-", p.path.casefold())
        named_file = any(re.search(r"(?:^|-)" + re.escape(token) + r"(?:-|$)", path_text) for token in tokens)
        # Manufacturer CDNs often place /in/en below /content/dam. Do not admit
        # documents from an explicitly different locale, even if the model matches.
        locale_in_path = not locale or any(segments[i:i+len(locale)] == locale for i in range(len(segments)))
        other_locale = _locale_prefix(url)
        if locale and other_locale and other_locale != locale:
            return False
        if locale and p.hostname == seed.hostname and not locale_in_path:
            return False
        return (named_file or named_label) and (linked or p.hostname == seed.hostname)
    if p.hostname != seed.hostname or (locale and segments[:len(locale)] != locale):
        return False
    # HTML must remain under the exact model segment. A label/slogan cannot
    # authorize a worldwide article, a similarly named model or another market.
    return model_segment


def ingest(demo_id: str, emit=lambda _: None, *, budget=None, fetcher=None) -> dict:
    """Discover and snapshot eligible sources. A later Read can continue deferred work.

    Serial requests stay below the approved two-per-host concurrency ceiling.
    Downloaded PDFs retain origin=website, so they never masquerade as uploads.
    """
    limits = {**DEFAULT_BUDGET, **(budget or {})}
    fetcher = fetcher or fetch_public
    demo, started = store.load(demo_id), time.monotonic()
    seeds = [s for s in demo.get("sources", []) if s.get("kind") == "url" and not s.get("crawl_parent") and s.get("use_in_demo", True)]
    prior = store.read_json(demo_id, "knowledge/coverage.json") or {}
    all_records, reports, excluded_prior_ids = [], [], set()
    totals = {"html_pages": 0, "documents": 0, "document_pages": 0}
    scoped_used, renders = {}, 0
    existing = {canonical_url(s["url"]): s for s in demo.get("sources", []) if s.get("url")}
    for seed in seeds:
        seed_url, tokens = canonical_url(seed["url"]), _model_tokens(seed, demo)
        scope_key = ("competition" if seed.get("role") == "competitor" else "product") + ":" + " ".join(tokens)
        used = scoped_used.setdefault(scope_key, {"html_pages": 0, "documents": 0, "document_pages": 0})
        for old_source in demo.get("sources", []):
            if old_source.get("crawl_parent") == seed["id"] and not _eligible(old_source.get("url", ""), old_source.get("name", ""), seed_url, tokens):
                excluded_prior_ids.add(old_source["id"])
        queue, seen, fetched, skipped, blocked, deferred = deque([(seed_url, 0, "", "seed")]), set(), [], [], [], []
        old = next((r for r in prior.get("seeds", []) if r.get("source_id") == seed["id"]), {})
        for item in old.get("deferred", []):
            if not _eligible(item["url"], "", seed_url, tokens):
                continue
            queue.append((item["url"], item.get("depth", 1), item.get("parent", seed_url), "continued"))
        root = f"{urlsplit(seed_url).scheme}://{urlsplit(seed_url).netloc}"
        robots, sitemap_urls, discovery_warnings = RobotFileParser(), [root + "/sitemap.xml"], []
        try:
            rob = fetcher(root + "/robots.txt", timeout=min(10, limits["seconds"]))
            robots.parse(rob.get("text", "").splitlines())
            sitemap_urls = robots.site_maps() or sitemap_urls
        except Exception as exc:
            robots.parse([])
            discovery_warnings.append(f"robots.txt unavailable: {str(exc)[:140]}")
        sitemap_seen = set()
        while sitemap_urls and len(sitemap_seen) < 8 and time.monotonic() - started < limits["seconds"]:
            sm = sitemap_urls.pop(0)
            if sm in sitemap_seen or urlsplit(sm).hostname != urlsplit(seed_url).hostname:
                continue
            sitemap_seen.add(sm)
            try:
                page = fetcher(sm, timeout=min(10, limits["seconds"]), max_bytes=3_000_000)
                xml = ElementTree.fromstring(page.get("raw") or page.get("text", ""))
                locs = [e.text.strip() for e in xml.iter() if e.tag.rsplit("}", 1)[-1] == "loc" and e.text]
                if xml.tag.rsplit("}", 1)[-1] == "sitemapindex":
                    sitemap_urls.extend(u for u in locs if not EXCLUDE.search(u))
                else:
                    if len(locs) > 20000:
                        discovery_warnings.append("Sitemap URL budget reached; locations after the first 20000 were not inspected.")
                    for u in locs[:20000]:
                        if _eligible(u, "", seed_url, tokens, linked=False):
                            queue.append((canonical_url(u), 1, sm, "sitemap"))
            except Exception as exc:
                discovery_warnings.append(f"Sitemap unavailable: {sm}: {str(exc)[:120]}")
        if sitemap_urls:
            discovery_warnings.append("Sitemap discovery budget reached; some sitemap indexes remain unread.")
        while queue:
            url, depth, parent, discovery = queue.popleft()
            if url in seen:
                continue
            seen.add(url)
            doc = bool(DOCUMENT.search(url))
            kind_budget = "documents" if doc else "html_pages"
            reason = ""
            if depth > limits["depth"]:
                reason = "depth budget"
            elif used[kind_budget] >= limits[kind_budget] or (doc and used["document_pages"] >= limits["document_pages"]):
                reason = "document/page budget" if doc else "HTML page budget"
            elif time.monotonic() - started >= limits["seconds"]:
                reason = "time budget"
            if reason:
                deferred.append({"url": url, "depth": depth, "parent": parent, "reason": reason})
                continue
            if urlsplit(url).hostname == urlsplit(seed_url).hostname and not robots.can_fetch(UA, url):
                blocked.append({"url": url, "reason": "robots.txt disallows this path"})
                continue
            try:
                emit(f"Reading source page {used['html_pages'] + used['documents'] + 1}: {url}")
                page = fetcher(url, timeout=min(20, max(1, limits["seconds"] - (time.monotonic() - started))))
                if page.get("error"):
                    raise ValueError(page["error"])
                final = canonical_url(page.get("final_url", url))
                if final != url and not _eligible(final, page.get("title", ""), seed_url, tokens):
                    raise ValueError("Redirect left the model-specific source scope")
                if not doc and renders < 4 and (len(page.get("text", "")) < 400 or
                        (b"<script" in page.get("raw", b"") and sum(len(s.get("text", "")) for s in page.get("sections", []) if s.get("kind") == "table") < 800 and "spec" in url.lower())):
                    renders += 1
                    try:
                        rendered = render_public(final, timeout=min(25, max(1, limits["seconds"] - (time.monotonic() - started))))
                        if len(rendered.get("text", "")) > len(page.get("text", "")):
                            page = rendered
                        else:
                            page.setdefault("warnings", []).append("Rendered fallback added no readable evidence.")
                    except Exception as exc:
                        page.setdefault("warnings", []).append(f"Rendered fallback unavailable or failed: {str(exc)[:160]}")
                is_pdf = page.get("mime") == "application/pdf" or final.lower().split("?")[0].endswith(".pdf")
                if is_pdf and not doc:
                    kind_budget = "documents"
                    if used[kind_budget] >= limits[kind_budget]:
                        deferred.append({"url": url, "depth": depth, "parent": parent, "reason": "document budget"})
                        continue
                used[kind_budget] += 1
                if is_pdf:
                    remaining = limits["document_pages"] - used["document_pages"]
                    count = page.get("page_count", len(page.get("pages", [])))
                    if count > remaining:
                        page["pages"] = page.get("pages", [])[:remaining]
                        page["sections"] = page.get("sections", [])[:remaining]
                        page["text"] = "\n\n".join(page["pages"])
                        page.setdefault("warnings", []).append("Document page budget reached; remaining pages require another bounded read.")
                    used["document_pages"] += min(len(page.get("pages", [])), remaining)
                raw = page.get("raw") or page.get("text", "").encode()
                digest = hashlib.sha256(raw).hexdigest()
                sid = seed["id"] if url == seed_url else existing.get(url, {}).get("id") or "src_" + hashlib.sha256(url.encode()).hexdigest()[:12]
                revision = f"knowledge/sources/{sid}/{digest}"
                raw_path = store.path(demo_id, revision + (".pdf" if is_pdf else ".html"))
                raw_path.parent.mkdir(parents=True, exist_ok=True)
                if not raw_path.exists():
                    raw_path.write_bytes(raw)
                if is_pdf and any(len(t.strip()) < 30 for t in page.get("pages", [])):
                    from .sources import ocr_pdf_pages
                    ocr_pages, ocr_report = ocr_pdf_pages(raw_path, page.get("pages", []), timeout=min(45, max(1, limits["seconds"] - (time.monotonic() - started))))
                    page["pages"] = ocr_pages
                    page["text"] = "\n\n".join(ocr_pages)
                    old_sections = page.get("sections", [])
                    page["sections"] = [{**(old_sections[i] if i < len(old_sections) else {}), "locator": f"page {i+1}", "kind": "pdf", "text": text} for i, text in enumerate(ocr_pages)]
                    page["ocr"] = ocr_report
                    if ocr_report:
                        page.setdefault("warnings", []).append("Local OCR attempted; OCR text and table column associations require review.")
                extracted = {k: v for k, v in page.items() if k != "raw"}
                extracted.update(id=sid, revision=digest, discovery_parent=parent, discovery=discovery, extraction_version=2)
                if not store.path(demo_id, revision + ".extract-v2.json").exists():
                    store.write_json(demo_id, revision + ".extract-v2.json", extracted)
                rec = {**existing.get(url, {}), "id": sid, "name": page.get("title") or url, "kind": "url", "url": url,
                       "role": seed.get("role", "product"), "origin": "website", "added_at": existing.get(url, {}).get("added_at", store.now()),
                       "use_in_demo": seed.get("use_in_demo", True), "crawl_active": True, "crawl_cached": False, "evidence_path": revision + ".extract-v2.json", "revision": digest,
                       "crawl_parent": seed.get("crawl_parent") if url == seed_url else seed["id"], "scope": seed.get("scope", {}),
                       "fetched_at": page.get("fetched_at", store.now()), "final_url": final, "extraction_version": 2}
                all_records.append(rec)
                fetched.append({"url": final, "source_id": sid, "revision": digest, "warnings": page.get("warnings", [])})
                for link in page.get("links", []):
                    if _eligible(link["url"], link.get("label", ""), seed_url, tokens):
                        queue.append((link["url"], depth + 1, final, "link"))
                    elif len(skipped) < 100:
                        skipped.append({"url": link["url"], "reason": "outside model scope"})
            except Exception as exc:
                blocked.append({"url": url, "reason": str(exc)[:300]})
        retained = [{"url": s["url"], "source_id": s["id"], "revision": s.get("revision", ""), "reason": "Earlier source revision retained outside this bounded read; verify freshness."}
                    for s in demo.get("sources", []) if s.get("crawl_parent") == seed["id"] and s["id"] not in excluded_prior_ids and s.get("evidence_path")
                    and canonical_url(s.get("url", "")) not in {p["url"] for p in fetched + blocked}]
        skipped.extend({"url": s.get("url", ""), "reason": "Earlier crawl excluded: outside exact model/market scope"} for s in demo.get("sources", []) if s["id"] in excluded_prior_ids and s.get("crawl_parent") == seed["id"])
        reports.append({"retained": retained, "source_id": seed["id"], "model_tokens": tokens, "fetched": fetched, "skipped": skipped, "blocked": blocked,
                        "deferred": deferred, "warnings": discovery_warnings + ([] if tokens else ["Model scope is ambiguous: only the supplied seed was eligible."]),
                        "frontier_exhausted": not retained and not deferred and not blocked and not discovery_warnings and not any(p.get("warnings") for p in fetched) and bool(tokens),
                        "topics_seen": [topic for topic in TOPICS if any(topic in p["url"].lower() for p in fetched)]})
    def merge(d):
        by_id = {s["id"]: s for s in all_records}
        seed_ids = {s["id"] for s in seeds}
        failed_urls = {canonical_url(p["url"]) for report in reports for p in report["blocked"]}
        def current(source):
            if source["id"] in excluded_prior_ids:
                return {**source, "crawl_active": False, "crawl_cached": False, "scope_excluded": True}
            if source["id"] in by_id:
                return by_id.pop(source["id"])
            if source["id"] in seed_ids or source.get("crawl_parent") in seed_ids:
                failed = canonical_url(source.get("url", "")) in failed_urls
                return {**source, "crawl_active": not failed and bool(source.get("evidence_path")), "crawl_cached": not failed}
            return source
        d["sources"] = [current(s) for s in d.get("sources", [])] + list(by_id.values())
    if seeds:
        store.update(demo_id, merge)
    coverage = {"version": 1, "at": store.now(), "budget": limits, "used": {key: sum(v[key] for v in scoped_used.values()) for key in totals}, "scopes": scoped_used, "seeds": reports,
                "complete": bool(reports) and all(r["frontier_exhausted"] for r in reports), "rendering": "local_chrome_playwright", "render_attempts": renders, "ocr": "local_tesseract" if shutil.which("tesseract") and shutil.which("pdftoppm") else "unavailable"}
    store.write_json(demo_id, "knowledge/coverage.json", coverage)
    return coverage
