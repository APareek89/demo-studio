"""Page-located source extraction, immutable revisions and loss-visible chunks."""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from . import config, store
from .crawl import UA, fetch_public, parse_html

EXTRACTION_VERSION = 2
PDF_IMAGE_VERSION = 1
PDF_IMAGE_LIMITS = {"max_pages": 100, "max_images": 96, "max_candidates": 256,
                    "max_file_bytes": 64 * 1024 * 1024, "max_image_bytes": 8 * 1024 * 1024,
                    "max_total_bytes": 48 * 1024 * 1024, "max_pixels": 12_000_000,
                    "max_edge": 2048, "timeout": 30}


def extract_pdf_images(p: Path, destination: Path, **overrides) -> dict:
    """Extract bounded embedded pictures, without interpreting text or inventing boxes.

    Raster data is decoded only after its PDF dimensions and encoded size pass
    limits. Normalized pixels determine identity; repeated logos reuse one asset.
    The caller retains the parent source/revision and applies ordinary vision.
    """
    from pypdf import PdfReader
    from PIL import Image
    limits = {**PDF_IMAGE_LIMITS, **overrides}
    started, images, warnings, seen = time.monotonic(), [], [], {}
    result = {"version": PDF_IMAGE_VERSION, "images": images, "warnings": warnings,
              "error": None, "interrupted": False, "candidates": 0, "bytes": 0}
    try:
        if p.stat().st_size > limits["max_file_bytes"]:
            raise ValueError("PDF exceeds the embedded-picture file-size budget")
        reader = PdfReader(p)
        if reader.is_encrypted:
            raise ValueError("encrypted PDF pictures need an unlocked source")
        candidates = []
        for page_number, page in enumerate(reader.pages, 1):
            if page_number > limits["max_pages"]:
                warnings.append("PDF picture page budget reached; remaining pages deferred.")
                break
            if time.monotonic() - started >= limits["timeout"]:
                result["interrupted"] = True
                warnings.append("PDF picture time budget reached; remaining pages deferred.")
                break
            for position, key in enumerate(page.images.keys()):
                if len(candidates) >= limits["max_candidates"]:
                    warnings.append("PDF picture candidate budget reached; remaining resources deferred.")
                    break
                candidates.append((position // 2, page_number, position, page, key))
            if len(candidates) >= limits["max_candidates"]:
                break
        # Give each page two opportunities before dense early galleries consume
        # the media budget. This is structural coverage, not a semantic claim.
        for _, page_number, position, page, key in sorted(candidates, key=lambda row: row[:3]):
            if time.monotonic() - started >= limits["timeout"]:
                result["interrupted"] = True
                warnings.append("PDF picture time budget reached; remaining pictures deferred.")
                return result
            if result["candidates"] >= limits["max_candidates"] or len(images) >= limits["max_images"]:
                warnings.append("PDF picture count budget reached; remaining pictures deferred.")
                return result
            result["candidates"] += 1
            try:
                # Resolve nested form resources before asking pypdf to decode.
                keys = [key] if isinstance(key, str) else key
                if len(keys) > 8 or any(str(name).startswith("~") for name in keys):
                    warnings.append(f"Page {page_number}: inline/deeply nested picture deferred.")
                    continue
                obj = page
                for name in keys:
                    obj = obj["/Resources"]["/XObject"][name].get_object()
                width, height = int(obj.get("/Width", 0)), int(obj.get("/Height", 0))
                if min(width, height) < 64:
                    continue  # Icons and fragments are not useful demo pictures.
                if width * height > limits["max_pixels"] or len(getattr(obj, "_data", b"")) > limits["max_image_bytes"]:
                    warnings.append(f"Page {page_number}: oversized picture deferred before decoding.")
                    continue
                mask = obj.get("/SMask")
                if mask:
                    mask = mask.get_object()
                    if (int(mask.get("/Width", 0)) * int(mask.get("/Height", 0)) > limits["max_pixels"]
                            or len(getattr(mask, "_data", b"")) > limits["max_image_bytes"]):
                        warnings.append(f"Page {page_number}: oversized picture mask deferred.")
                        continue
                encoded = page.images[key].data
                if len(encoded) > limits["max_image_bytes"]:
                    warnings.append(f"Page {page_number}: decoded picture exceeds byte budget.")
                    continue
                with Image.open(io.BytesIO(encoded)) as original:
                    if original.width * original.height > limits["max_pixels"]:
                        warnings.append(f"Page {page_number}: decoded picture exceeds pixel budget.")
                        continue
                    image = original.convert("RGBA" if "A" in original.getbands() or "transparency" in original.info else "RGB")
                    image.thumbnail((limits["max_edge"], limits["max_edge"]))
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    payload = buffer.getvalue()
                    image_size = list(image.size)
                digest = hashlib.sha256(payload).hexdigest()
                occurrence = {"page": page_number, "image_index": position + 1, "resource": key}
                if digest in seen:
                    if page_number not in seen[digest]["pages"]:
                        seen[digest]["pages"].append(page_number)
                    seen[digest]["occurrences"].append(occurrence)
                    continue
                if len(payload) > limits["max_image_bytes"] or result["bytes"] + len(payload) > limits["max_total_bytes"]:
                    warnings.append(f"Page {page_number}: normalized picture exceeds remaining byte budget.")
                    continue
                destination.mkdir(parents=True, exist_ok=True)
                path = destination / f"{digest}.png"
                if not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                    with tempfile.NamedTemporaryFile(dir=destination, suffix=".tmp", delete=False) as pending:
                        pending.write(payload)
                        pending_path = Path(pending.name)
                    try:
                        pending_path.replace(path)
                    finally:
                        pending_path.unlink(missing_ok=True)
                row = {"path": path.name, "sha256": digest, "pages": [page_number],
                       "occurrences": [occurrence], "size": len(payload), "dimensions": image_size}
                images.append(row)
                seen[digest] = row
                result["bytes"] += len(payload)
            except Exception as exc:
                warnings.append(f"Page {page_number}: picture unavailable ({str(exc)[:120]}).")
        return result
    except Exception as exc:
        result["error"] = f"could not extract PDF pictures: {str(exc)[:180]}"
        return result


def ocr_pdf_pages(p: Path, pages: list[str], *, max_ocr_pages: int = 12, timeout: float = 90) -> tuple[list[str], list[dict]]:
    """Use installed local Poppler/Tesseract only; no model spend or new dependency."""
    renderer, ocr = shutil.which("pdftoppm"), shutil.which("tesseract")
    report, out, started = [], list(pages), time.monotonic()
    sparse = [i for i, text in enumerate(pages) if len(text.strip()) < 30]
    for i in sparse:
        if not renderer or not ocr:
            report.append({"page": i+1, "status": "unavailable", "reason": "pdftoppm/tesseract not installed"})
            continue
        if len([r for r in report if r["status"] != "deferred"]) >= max_ocr_pages or time.monotonic() - started >= timeout:
            report.append({"page": i+1, "status": "deferred", "reason": "OCR page/time budget"})
            continue
        try:
            with tempfile.TemporaryDirectory(prefix="demo-ocr-") as directory:
                prefix = str(Path(directory) / "page")
                remaining = max(1, timeout - (time.monotonic() - started))
                subprocess.run([renderer, "-f", str(i+1), "-l", str(i+1), "-r", "150", "-singlefile", "-png", str(p), prefix],
                               check=True, capture_output=True, timeout=min(30, remaining))
                remaining = max(1, timeout - (time.monotonic() - started))
                result = subprocess.run([ocr, prefix + ".png", "stdout", "--psm", "6"], check=True, capture_output=True,
                                        text=True, timeout=min(30, remaining))
                out[i] = result.stdout.strip()
                report.append({"page": i+1, "status": "extracted" if out[i] else "empty", "method": "local_tesseract", "review_required": True})
        except (OSError, subprocess.SubprocessError) as exc:
            report.append({"page": i+1, "status": "failed", "reason": str(exc)[:150]})
    return out, report


def extract_pdf(p, *, max_pages: int = 300, timeout: float = 90) -> dict:
    """Retain measured line alignment and raw table cells, never fill merged cells.

    Extraction versions are separate from immutable source-byte revisions. This
    avoids pypdf layout's pathological fixed-pitch padding in automotive brochures.
    """
    import pdfplumber
    started, pages, sections, warnings = time.monotonic(), [], [], []
    try:
        with pdfplumber.open(p) as document:
            count = len(document.pages)
            for i, page in enumerate(document.pages[:max_pages]):
                if time.monotonic() - started > timeout:
                    warnings.append("PDF extraction time budget reached; remaining pages were deferred.")
                    break
                raw = page.extract_text(layout=True) or ""
                text = "\n".join(line.rstrip() for line in raw.splitlines() if line.strip())
                # Whitespace is layout context, not millions of prompt characters.
                if len(text) > 60000 or any(len(line) > 1500 for line in text.splitlines()):
                    text = page.extract_text(layout=False) or ""
                    warnings.append(f"Page {i+1}: dense text fallback used; verify table associations in the original.")
                tables = page.extract_tables() or []
                pages.append(text)
                sections.append({"locator": f"page {i+1}", "kind": "pdf", "text": text, "tables": tables,
                                 "extraction": "pdfplumber-layout", "table_note": "Null/empty cells remain unresolved; never infer merged headers or missing values." if tables else ""})
                page.close()
            if count > len(pages):
                warnings.append(f"PDF page/time budget reached: {len(pages)} of {count} pages extracted.")
        return {"pages": pages, "sections": sections, "page_count": count, "warnings": warnings, "error": None}
    except Exception as exc:
        return {"pages": [], "sections": [], "page_count": 0, "warnings": [], "error": f"could not read pdf: {exc}"}


def pdf_pages(p: Path, *, max_pages: int = 300) -> list[str]:
    result = extract_pdf(p, max_pages=max_pages)
    return result["pages"] if not result["error"] else [f"[{result['error']}]"]


def docx_text(p: Path) -> str:
    import docx
    try:
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in table.rows))
        return "\n\n".join(x for x in parts if x.strip())
    except Exception as e:
        return f"[could not read docx: {e}]"


def fetch_url(url: str) -> dict:
    """Backwards-compatible one-page fetch; discovery belongs to crawl.ingest."""
    if config.MOCK_LLM:
        return {"title": "", "text": "Mock source; no network fetch.", "images": [], "warnings": ["Mock: URL not fetched."], "error": None}
    try:
        return {k: v for k, v in fetch_public(url, timeout=25).items() if k != "raw"}
    except Exception as e:
        return {"url": url, "title": "", "text": "", "images": [], "sections": [], "warnings": [], "error": str(e)}


def chunks(extraction: dict, *, max_chars: int = 24000) -> list[dict]:
    """No silent source truncation. Long tables repeat their header and footnotes."""
    sections = extraction.get("sections") or [{"locator": "document", "kind": "text", "text": extraction.get("text", "")}]
    result = []
    for section in sections:
        text = section.get("text", "")
        pieces = []
        if len(text) <= max_chars:
            pieces = [text]
        elif section.get("rows"):
            rows = section["rows"]
            header = " | ".join(rows[0]) if rows else ""
            suffix = "\nFootnote: " + section["footnote"] if section.get("footnote") else ""
            current = header
            for row in rows[1:]:
                line = " | ".join(row)
                available = max(100, max_chars - len(header) - len(suffix) - 1)
                if len(line) > available:
                    if current != header:
                        pieces.append(current + suffix)
                        current = header
                    for offset in range(0, len(line), available):
                        pieces.append(header + "\n" + line[offset:offset+available] + suffix)
                    continue
                if len(current) + len(line) + len(suffix) > max_chars and current != header:
                    pieces.append(current + suffix)
                    current = header
                current += "\n" + line
            pieces.append(current + suffix)
        else:
            # Split on line boundaries where possible, including a little exact context.
            while len(text) > max_chars:
                split = text.rfind("\n", 0, max_chars)
                split = split if split > max_chars // 2 else max_chars
                pieces.append(text[:split])
                text = text[max(0, split - 160):]
            if text:
                pieces.append(text)
        for n, piece in enumerate(pieces):
            if not piece.strip():
                continue
            result.append({"id": "chunk_" + hashlib.sha256((str(extraction.get("id")) + section.get("locator", "") + piece).encode()).hexdigest()[:20],
                           "source_id": extraction.get("id"), "locator": section.get("locator", "document"), "part": n + 1,
                           "kind": section.get("kind", "text"), "text": piece, "tables": section.get("tables", []), "table_note": section.get("table_note", "")})
    return result


def source_text(demo_id: str, src: dict, *, persist: bool = False, max_pages: int = 300) -> dict:
    """Extract every eligible chunk, preserving existing consumer keys.

    persist is only set during Read; ordinary callers never change source metadata.
    Crawled URL revisions are already cached and never fetched again here.
    """
    kind = src["kind"]
    cached = store.read_json(demo_id, src["evidence_path"]) if src.get("evidence_path") else None
    if kind == "url" and cached and cached.get("extraction_version", 1) < EXTRACTION_VERSION:
        # Reparse original immutable bytes. Never label an old parse as current,
        # and never overwrite the prior extraction used by a published snapshot.
        base = f"knowledge/sources/{src['id']}/{src.get('revision', '')}"
        raw_file = next((store.path(demo_id, base + suffix) for suffix in (".pdf", ".html", ".bin") if store.path(demo_id, base + suffix).exists()), None)
        if raw_file and raw_file.suffix == ".pdf":
            pdf = extract_pdf(raw_file, max_pages=max_pages)
            cached = {**cached, **pdf, "text": "\n\n".join(pdf["pages"]), "extraction_version": EXTRACTION_VERSION}
        elif raw_file and "html" in cached.get("mime", "text/html"):
            cached = {**cached, **parse_html(raw_file.read_bytes(), cached.get("final_url") or src["url"]), "extraction_version": EXTRACTION_VERSION}
        else:
            # No trustworthy raw revision: a safe fetch must establish fresh evidence.
            cached = None
    if kind == "url":
        f = cached or fetch_url(src["url"])
        out = {**f, "id": src["id"], "name": src["url"], "kind": "url", "pages": f.get("pages"), "images": f.get("images", []),
               "text": f.get("text", "") if not f.get("error") else f"[fetch failed: {f['error']}]"}
        raw = out.get("text", "").encode()
    else:
        p = store.path(demo_id, src["path"])
        try:
            raw = p.read_bytes()
        except OSError as exc:
            raw = b""
            out = {"error": str(exc), "text": "", "pages": None}
        else:
            warnings, pages = [], None
            if kind == "pdf":
                pdf = extract_pdf(p, max_pages=max_pages)
                pages, warnings = pdf["pages"], pdf["warnings"]
                if pdf["error"]:
                    warnings.append(pdf["error"])
                pages, ocr_report = ocr_pdf_pages(p, pages)
                if ocr_report:
                    warnings.append("OCR fallback: " + "; ".join(f"page {r['page']} {r['status']}" for r in ocr_report) + ". OCR text requires source review; column associations are not certified.")
                if any(len(page.strip()) < 30 for page in pages):
                    warnings.append("Some pages still contain little text after bounded local OCR. Coverage remains incomplete.")
                sections = [{**section, "text": pages[i]} for i, section in enumerate(pdf["sections"])]
                text = "\n\n".join(f"[page {i+1}]\n{t}" for i, t in enumerate(pages))
            elif kind == "doc":
                text = docx_text(p)
                sections = [{"locator": "document", "kind": "text", "text": text}]
                if p.suffix.lower() == ".doc":
                    warnings.append("Legacy .doc extraction is unsupported; upload .docx or PDF.")
            else:
                text = raw.decode("utf-8", errors="replace") if kind == "text" else ""
                sections = [{"locator": "document", "kind": "text", "text": text}]
            out = {"text": text, "pages": pages, "sections": sections, "warnings": warnings, "error": None, "ocr": ocr_report if kind == "pdf" else []}
            if text.startswith("[could not read"):
                out["error"] = text
        out.update(id=src["id"], name=src["name"], kind=kind)
    revision = (src.get("revision") if cached else out.get("hash")) if kind == "url" else None
    revision = revision or hashlib.sha256(raw).hexdigest()
    out["revision"] = revision
    out["extraction_version"] = EXTRACTION_VERSION
    if src.get("crawl_cached"):
        out.setdefault("warnings", []).append("Earlier source revision reused outside the current bounded read; verify freshness before approving.")
    out["chunks"] = chunks(out)
    if persist:
        evidence_path = f"knowledge/sources/{src['id']}/{revision}.extract-v{EXTRACTION_VERSION}.json"
        if not cached or kind != "url":
            raw_path = store.path(demo_id, f"knowledge/sources/{src['id']}/{revision}.bin")
            raw_path.parent.mkdir(parents=True, exist_ok=True)
            if not raw_path.exists():
                raw_path.write_bytes(raw)
        if not store.path(demo_id, evidence_path).exists():
            store.write_json(demo_id, evidence_path, out)
        def update(d):
            source = next(s for s in d["sources"] if s["id"] == src["id"])
            source.update(revision=revision, evidence_path=evidence_path, extraction_version=EXTRACTION_VERSION, origin=source.get("origin", "website" if kind == "url" else "uploaded"))
        store.update(demo_id, update)
    return out


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40]
